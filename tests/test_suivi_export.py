"""Notes -> Suivi's Journal sheet. The harness's first write to one of her own files, so the tests are strict."""
import json
import os
import re
import sqlite3
import time
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from harness import db
from harness import suivi_export as SE
from tools import suivi_writer as W

HEADER = ["id", "date", "domain", "code", "type", "summary", "people", "evidence", "hours", "source", "status", "entered_on"]
NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'


def _row(r, values, styles):
    cells = []
    for col, v, s in zip(W.COLS, values, styles):
        if v is None:
            cells.append(f'<c r="{col}{r}" s="{s}" t="n"/>')
        else:
            cells.append(f'<c r="{col}{r}" s="{s}" t="inlineStr"><is><t>{v}</t></is></c>')
    return f'<row r="{r}" ht="15.75" customHeight="1" s="7">' + "".join(cells) + "</row>"


def make_suivi(path: Path, entries: list[list[str]], total_rows: int = 30) -> None:
    """A workbook shaped like hers: inline strings, header row, entries, then empty pre-formatted rows; autofilter and dropdowns after the data."""
    st = [0, 0, 6, 0, 6, 0, 0, 0, 0, 6, 6, 0]
    rows = ['<row r="1">' + "".join(f'<c r="{c}1" s="1" t="inlineStr"><is><t>{h}</t></is></c>' for c, h in zip(W.COLS, HEADER)) + "</row>"]
    for i, e in enumerate(entries, start=2):
        rows.append(_row(i, e, st))
    for k in range(len(entries) + 2, total_rows + 1):
        rows.append(f'<row r="{k}" ht="15.75" customHeight="1" s="7"><c r="C{k}" s="6" t="n"/><c r="E{k}" s="6" t="n"/><c r="J{k}" s="6" t="n"/><c r="K{k}" s="6" t="n"/></row>')
    sheet1 = (f'<worksheet {NS}><sheetPr><outlinePr summaryBelow="1" summaryRight="1"/></sheetPr><dimension ref="A1:L{total_rows}"/>'
              f'<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
              f'<sheetFormatPr baseColWidth="8" defaultColWidth="14.43" defaultRowHeight="15" customHeight="1"/><cols><col width="11" customWidth="1" style="7" min="1" max="2"/></cols>'
              f'<sheetData>{"".join(rows)}</sheetData><autoFilter ref="A1:L5"/><dataValidations count="1"><dataValidation sqref="E2:E{total_rows}" type="list" allowBlank="1">'
              f'<formula1>"meeting,call,email,work,decision,event,admin,document"</formula1></dataValidation></dataValidations></worksheet>')
    sheet2 = f'<worksheet {NS}><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>id</t></is></c></row><row r="2"><c r="A2" t="inlineStr"><is><t>C-2026-001</t></is></c></row></sheetData></worksheet>'
    styles = (f'<styleSheet {NS}><fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts><fills count="2"><fill><patternFill patternType="none"/></fill>'
              '<fill><patternFill patternType="gray125"/></fill></fills><borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
              '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="8">' + '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>' * 8 + '</cellXfs></styleSheet>')
    parts = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>',
        "_rels/.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook {NS} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Journal" sheetId="1" r:id="rId1"/><sheet name="Commitments" sheetId="2" r:id="rId2"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet2.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="/xl/styles.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": sheet1, "xl/worksheets/sheet2.xml": sheet2, "xl/styles.xml": styles,
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    os.utime(path, (time.time() - 3600, time.time() - 3600))                     # an hour old: "not being saved right now"


def entry(n, summary="Earlier entry"):
    return [f"J-2026-{n:03d}", "2026-09-2" + str(n % 10), "W", "GPF", "meeting", summary, "dan", "cal:abc", None, "capture", "confirmed", "2026-09-25"]


@pytest.fixture
def suivi(tmp_path):
    p = tmp_path / "Drive" / "Suivi.xlsx"; p.parent.mkdir()
    make_suivi(p, [entry(1), entry(2), entry(3), entry(4)])
    return p


def rows_of(path, sheet="Journal"):
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    return [[c if c is not None else "" for c in r] for r in ws.iter_rows(values_only=True)]            # an empty cell reads as ""


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


NEW = [["J-2026-005", "2026-10-02", "W", "GPF", "work", "Call the venue about the projector", "", "harness-note:1", "", "capture", "confirmed", "2026-10-02"],
       ["J-2026-006", "2026-10-02", "W", "", "email", "Spoke to Daniel: he prefers Thursday", "dan", "18c0ffee12345678", "", "capture", "confirmed", "2026-10-02"]]


# ---------------------------------------------------------------- the append itself
def test_new_rows_land_in_the_next_empty_rows_and_nothing_else_changes(suivi, tmp_path):
    before_parts, before_rows = parts(suivi), rows_of(suivi)
    used = W.append_to_suivi(suivi, NEW, backups=tmp_path / "bk")
    assert used == [6, 7]
    after = rows_of(suivi)
    assert after[:5] == before_rows[:5]                                                   # header and the four entries, untouched
    assert after[5][:8] == NEW[0][:8] and after[6][4] == "email" and after[6][7] == "18c0ffee12345678"
    assert after[7:] == before_rows[7:]                                                   # everything below stays empty
    now_parts = parts(suivi)
    assert [n for n in now_parts if now_parts[n] != before_parts[n]] == ["xl/worksheets/sheet1.xml"]
    assert rows_of(suivi, "Commitments") == [["id"], ["C-2026-001"]]
    xml = now_parts["xl/worksheets/sheet1.xml"].decode()
    for kept in ('<dimension ref="A1:L30"/>', '<autoFilter ref="A1:L5"/>', '<dataValidations count="1">', 'state="frozen"', 'customHeight="1" s="7"'):
        assert kept in xml


def test_styles_follow_the_sheets_own_formatting(suivi, tmp_path):
    W.append_to_suivi(suivi, NEW[:1], backups=tmp_path / "bk")
    cells = W.cells_of(W.scan_rows(parts(suivi)["xl/worksheets/sheet1.xml"].decode())[5]["inner"])
    assert [cells[c]["s"] for c in "ABCDEFGHIJKL"] == ["0", "0", "6", "0", "6", "0", "0", "0", "0", "6", "6", "0"]


def test_awkward_characters_survive_the_round_trip(suivi, tmp_path):
    tricky = "Q&A <draft> “quoted” it's 5 > 3 — é ü 日本語 🎉 \x07bell"
    row = [NEW[0][0], "2026-10-02", "W", "", "work", W_clean(tricky), "", "harness-note:2", "", "capture", "confirmed", "2026-10-02"]
    W.append_to_suivi(suivi, [row], backups=tmp_path / "bk")
    got = rows_of(suivi)[5][5]
    assert got == "Q&A <draft> “quoted” it's 5 > 3 — é ü 日本語 🎉 bell"                   # the control character is dropped, the rest intact


def W_clean(s):
    return SE.clean_summary(s)


def test_a_backup_is_made_before_the_write_and_old_ones_are_pruned(suivi, tmp_path):
    bk = tmp_path / "bk"; bk.mkdir()
    for i in range(35):
        (bk / f"Suivi-20260101-{i:06d}.xlsx").write_text("old")
    original = suivi.read_bytes()
    W.append_to_suivi(suivi, NEW, backups=bk)
    kept = sorted(bk.glob("Suivi-*.xlsx"))
    assert len(kept) == W.KEEP_BACKUPS and kept[-1].read_bytes() == original and not (suivi.parent / ".Suivi.xlsx.harness-tmp").exists()


def test_nothing_to_add_does_nothing(suivi, tmp_path):
    before = suivi.read_bytes()
    assert W.append_to_suivi(suivi, [], backups=tmp_path / "bk") == [] and suivi.read_bytes() == before


# ---------------------------------------------------------------- when it must refuse
def test_it_waits_while_excel_has_the_file_open(suivi, tmp_path):
    before = suivi.read_bytes()
    (suivi.parent / "~$Suivi.xlsx").write_text("lock")
    with pytest.raises(W.SuiviBusy, match="open in a spreadsheet"):
        W.append_to_suivi(suivi, NEW, backups=tmp_path / "bk")
    assert suivi.read_bytes() == before
    (suivi.parent / "~$Suivi.xlsx").unlink(); (suivi.parent / ".~lock.Suivi.xlsx#").write_text("lo")
    with pytest.raises(W.SuiviBusy):
        W.append_to_suivi(suivi, NEW, backups=tmp_path / "bk")


def test_it_waits_if_the_file_changed_in_the_last_minute_and_a_half(suivi, tmp_path):
    os.utime(suivi, None)
    with pytest.raises(W.SuiviBusy, match="syncing or being saved"):
        W.append_to_suivi(suivi, NEW, backups=tmp_path / "bk")


def test_it_aborts_if_the_file_changes_between_the_final_check_and_the_swap(suivi, tmp_path, monkeypatch):
    real = W.verify_result
    state = {"first": True}
    def verify_then_she_saves(*a, **k):
        real(*a, **k)
        if state["first"]:                                                                 # the new file passed its check; now SHE saves a change
            state["first"] = False
            with zipfile.ZipFile(suivi, "a") as z:
                z.writestr("customXml/note.xml", "<x/>")
            os.utime(suivi, (time.time() - 3600, time.time() - 3600))
    monkeypatch.setattr(W, "verify_result", verify_then_she_saves)
    with pytest.raises(W.SuiviBusy, match="changed while preparing"):
        W.append_to_suivi(suivi, NEW, backups=tmp_path / "bk")
    assert b"customXml/note.xml" in suivi.read_bytes() and not (suivi.parent / ".Suivi.xlsx.harness-tmp").exists()      # her change is intact, ours was dropped
    assert not (tmp_path / "bk").exists() or list((tmp_path / "bk").glob("Suivi-*.xlsx")) == []


def test_it_refuses_when_the_sheet_has_no_empty_rows_left(tmp_path):
    p = tmp_path / "Suivi.xlsx"; make_suivi(p, [entry(i) for i in range(1, 5)], total_rows=6)       # rows 6 is the last: only one slot
    with pytest.raises(W.SuiviError, match="no empty rows"):
        W.append_to_suivi(p, NEW, backups=tmp_path / "bk")


def test_it_refuses_a_sheet_it_does_not_recognise(tmp_path):
    p = tmp_path / "Suivi.xlsx"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("xl/workbook.xml", '<workbook><sheets><sheet name="Other" r:id="rId1"/></sheets></workbook>'); z.writestr("xl/_rels/workbook.xml.rels", "<Relationships/>")
    os.utime(p, (time.time() - 3600, time.time() - 3600))
    with pytest.raises(W.SuiviError, match="no sheet named Journal"):
        W.append_to_suivi(p, NEW, backups=tmp_path / "bk")


# ---------------------------------------------------------------- safety nets
def test_if_the_result_is_not_sound_after_the_swap_her_original_is_put_back(suivi, tmp_path):
    original = suivi.read_bytes()
    def damage(path):                                                                      # something goes wrong right after the swap
        with zipfile.ZipFile(path, "a") as z:
            z.writestr("xl/extra.xml", "<oops/>")
    with pytest.raises(W.SuiviError):
        W.append_to_suivi(suivi, NEW, backups=tmp_path / "bk", _sabotage=damage)
    assert suivi.read_bytes() == original                                                  # byte for byte
    assert len(list((tmp_path / "bk").glob("Suivi-*.xlsx"))) == 1


def test_verification_catches_a_change_anywhere_it_should_not_be(suivi, tmp_path):
    xml = parts(suivi)["xl/worksheets/sheet1.xml"].decode()
    new_xml, used, _ = W.fill_rows(xml, NEW)
    bad = tmp_path / "bad.xlsx"
    W.rewrite_zip(suivi, bad, "xl/worksheets/sheet1.xml", new_xml.replace("Earlier entry", "Tampered entry", 1))
    with pytest.raises(W.SuiviError, match="outside the new rows"):
        W.verify_result(suivi, bad, "xl/worksheets/sheet1.xml", used, NEW)
    other = tmp_path / "other.xlsx"
    with zipfile.ZipFile(suivi) as a, zipfile.ZipFile(other, "w") as z:
        for n in a.namelist():
            z.writestr(n, new_xml if n.endswith("sheet1.xml") else (b"<changed/>" if n.endswith("sheet2.xml") else a.read(n)))
    with pytest.raises(W.SuiviError, match="sheet2.xml changed"):
        W.verify_result(suivi, other, "xl/worksheets/sheet1.xml", used, NEW)
    wrong = tmp_path / "wrong.xlsx"
    W.rewrite_zip(suivi, wrong, "xl/worksheets/sheet1.xml", new_xml.replace("Call the venue", "Call someone else", 1))
    with pytest.raises(W.SuiviError, match="does not contain"):
        W.verify_result(suivi, wrong, "xl/worksheets/sheet1.xml", used, NEW)


# ---------------------------------------------------------------- from notes to rows
@pytest.fixture
def world(tmp_path):
    data = tmp_path / "data"; data.mkdir()
    c = db.connect(data / "harness.db"); db.migrate(c)
    c.execute("INSERT INTO tasks (title, due_date, project_code, source, source_ref) VALUES ('Prepare panel', '2026-10-05', 'GPF', 'suivi', 'C-2026-017')")
    c.execute("INSERT INTO emails (thread_id, message_id, from_email, subject, received_at, person_slug) VALUES ('18c0ffee12345678','m1','dan@org.ch','Hello','2026-10-01T09:00:00+02:00','dan')")
    c.execute("INSERT INTO waiting_on (description, person, since_date) VALUES ('the agreement','nadja','2026-09-20')")
    def note(text, kind="note", pt=None, pid=None, due=None, space="work", deleted=None, code=None):
        c.execute("INSERT INTO notes (created_at, text, kind, parent_type, parent_id, due_date, project_code, space, deleted_at) VALUES ('2026-10-02 09:30:00',?,?,?,?,?,?,?,?)", (text, kind, pt, pid, due, code, space, deleted))
    note("Free-standing idea about the forum")                                                  # 1
    note("Spoke to Daniel: Thursday works.\nWill confirm.", pt="email", pid=1)                  # 2
    note("Call the venue", kind="followup", pt="task", pid=1, due="2026-10-08", code="GPF")     # 3
    note("Private matter", space="personal")                                                    # 4: never exported
    note("Changed my mind", deleted="2026-10-02 10:00:00")                                      # 5: deleted
    note("Chased by phone", pt="waiting_on", pid=1)                                             # 6
    c.commit(); c.close()
    suivi = tmp_path / "Drive"; suivi.mkdir(); make_suivi(suivi / "Suivi.xlsx", [entry(i) for i in range(1, 6)] + [entry(7)])      # highest id J-2026-007 (006 is a gap)
    return data, suivi


def test_work_notes_become_journal_rows_following_the_spec(world):
    data, suivi = world
    out = W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    assert out == "added 4 note(s) to Suivi: J-2026-008, J-2026-009, J-2026-010, J-2026-011"        # not personal, not deleted; the gap at 006 is not reused
    rows = {r[0]: r for r in rows_of(suivi / "Suivi.xlsx") if r[0] and r[0] >= "J-2026-008"}
    free, email, follow, chased = rows["J-2026-008"], rows["J-2026-009"], rows["J-2026-010"], rows["J-2026-011"]
    assert free[1:6] == ["2026-10-02", "W", "", "work", "Free-standing idea about the forum"] and free[7] == "harness-note:1" and free[9:12] == ["capture", "confirmed", "2026-10-02"]
    assert email[4] == "email" and email[5] == "Spoke to Daniel: Thursday works. / Will confirm." and email[6] == "dan" and email[7] == "18c0ffee12345678"      # evidence is the Gmail thread ID
    assert follow[5] == "Follow-up (due 2026-10-08): Call the venue" and follow[3] == "GPF" and follow[7] == "harness-note:3; C-2026-017"
    assert chased[6] == "nadja"
    state = SE.read_state(data)
    assert state["notes"]["2"]["journal_id"] == "J-2026-009" and set(state["notes"]) == {"1", "2", "3", "6"} and state["last_error"] is None


def test_a_note_is_exported_once_only(world):
    data, suivi = world
    W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    snapshot = (suivi / "Suivi.xlsx").read_bytes()
    assert W.run(data=data, suivi=suivi, today=date(2026, 10, 2)) == "nothing to add" and (suivi / "Suivi.xlsx").read_bytes() == snapshot
    c = db.connect(data / "harness.db")
    c.execute("INSERT INTO notes (created_at, text, space) VALUES ('2026-10-03 08:00:00','A later note','work')"); c.commit(); c.close()
    os.utime(suivi / "Suivi.xlsx", (time.time() - 3600, time.time() - 3600))
    assert "J-2026-012" in W.run(data=data, suivi=suivi, today=date(2026, 10, 3))


def test_dry_run_shows_the_rows_and_changes_nothing(world):
    data, suivi = world
    before = (suivi / "Suivi.xlsx").read_bytes()
    out = W.run(True, data=data, suivi=suivi, today=date(2026, 10, 2))
    assert out.startswith("would add:") and "Free-standing idea" in out and "Private matter" not in out
    assert (suivi / "Suivi.xlsx").read_bytes() == before and SE.read_state(data)["notes"] == {}


def test_the_pause_switch_stops_everything(world):
    data, suivi = world
    (data / SE.STATE_DIR).mkdir(); (data / SE.STATE_DIR / SE.DISABLED_FILE).write_text("")
    before = (suivi / "Suivi.xlsx").read_bytes()
    assert W.run(data=data, suivi=suivi).startswith("paused") and (suivi / "Suivi.xlsx").read_bytes() == before


def test_if_suivi_is_busy_the_notes_stay_pending(world):
    data, suivi = world
    (suivi / "~$Suivi.xlsx").write_text("lock")
    with pytest.raises(W.SuiviBusy):
        W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    assert SE.read_state(data)["notes"] == {}


def test_the_existing_importer_reads_the_new_rows_back_as_journal_entries(world):
    from harness.importers.suivi import import_suivi
    data, suivi = world
    W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    # the importer needs the sheets it knows; add empty ones next to our stand-in by reading only the Journal part
    c = db.connect(data / "harness.db")
    path = suivi / "Suivi.xlsx"
    wb = openpyxl.load_workbook(path)
    for name, header in (("Commitments", ["id", "created", "domain", "code", "direction", "counterparty", "what", "due", "weight", "status", "closed_on"]),
                         ("People", ["id", "name", "aliases", "org", "role", "email", "cadence", "domain"])):
        if name in wb.sheetnames:
            del wb[name]
        wb.create_sheet(name).append(header)
    wb.save(path)
    reports = import_suivi(c, suivi)
    journal = next(r for r in reports if r.source == "suivi:journal")
    assert journal.notes == [] and journal.added >= 4
    row = c.execute("SELECT entry_date, project_code, text FROM journal_entries WHERE source_ref='J-2026-010'").fetchone()
    assert tuple(row) == ("2026-10-02", "GPF", "Follow-up (due 2026-10-08): Call the venue")


# ---------------------------------------------------------------- the pure rules
def test_next_journal_id_continues_from_the_highest_and_ignores_other_years_and_junk():
    assert SE.next_journal_id([], 2026) == "J-2026-001"
    assert SE.next_journal_id(["J-2026-001", "J-2026-007", "J-2026-003"], 2026) == "J-2026-008"
    assert SE.next_journal_id(["J-2025-099", "J-2026-002", "x", "", None, "J-2026-1x"], 2026) == "J-2026-003"
    assert SE.next_journal_id(["J-2026-999"], 2026) == "J-2026-1000"
    assert SE.next_journal_id(["J-2026-005"], 2027) == "J-2027-001"


def test_summary_is_one_line_and_bounded():
    assert SE.clean_summary("  one\n\n two \n three ") == "one / two / three"
    long = SE.clean_summary("x" * 2000)
    assert len(long) == SE.SUMMARY_LIMIT and long.endswith("…")
    assert SE.clean_summary("Call", followup_due="2026-10-08") == "Follow-up (due 2026-10-08): Call"


def test_note_status_for_the_screen():
    st = {"notes": {"7": {"journal_id": "J-2026-012"}}}
    assert SE.note_status(7, "work", False, st) == {"state": "exported", "journal_id": "J-2026-012"}
    assert SE.note_status(8, "work", False, st) == {"state": "pending"} and SE.note_status(8, "work", False, st, paused=True) == {"state": "paused"}
    assert SE.note_status(9, "personal", False, st)["state"] == "excluded" and SE.note_status(9, "work", True, st) == {"state": "none"}


def test_only_work_notes_are_eligible():
    assert SE.eligible("work", False) and not SE.eligible("personal", False) and not SE.eligible("work", True)


# ---------------------------------------------------------------- filed documents -> journal rows
def _filed(data, **kw):
    c = db.connect(data / "harness.db")
    row = dict(source="folder", original_name="Scan.pdf", sha256="a" * 64, size=1, status="filed", doc_type="invoice_received", supplier="Alber Rolle SA", number="002131", amount=2972.75,
               currency="CHF", paid_date="2026-01-29", folder="Expenses", proposed_name="Alber_Rolle_SA_Facture_002131_CHF2972.75_paye_29.01.2026.pdf", year="2026", filed_at="2026-10-02 10:15:00")
    row.update(kw)
    c.execute(f"INSERT INTO scans ({','.join(row)}) VALUES ({','.join('?' * len(row))})", list(row.values()))
    c.commit(); c.close()


def test_a_filed_document_becomes_a_journal_row_with_the_file_as_evidence(world):
    data, suivi = world
    _filed(data)
    out = W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    assert out == "added 4 note(s) and 1 document(s) to Suivi: J-2026-008, J-2026-009, J-2026-010, J-2026-011, J-2026-012"
    row = next(r for r in rows_of(suivi / "Suivi.xlsx") if r[0] == "J-2026-012")
    assert row[1:6] == ["2026-10-02", "W", "FIN", "document", "Filed invoice: Alber Rolle SA, no. 002131, CHF 2972.75, paid 29.01.2026 (in 2026/Expenses)"]
    assert row[7] == "Admin/ICT4Peace Audit/2026/Expenses/Alber_Rolle_SA_Facture_002131_CHF2972.75_paye_29.01.2026.pdf" and row[9:12] == ["capture", "confirmed", "2026-10-02"]
    assert SE.read_state(data)["docs"]["1"]["journal_id"] == "J-2026-012"


def test_a_document_is_added_once_and_only_when_it_is_really_filed(world):
    data, suivi = world
    W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    _filed(data, status="approved", sha256="b" * 64)                      # approved but not yet written by the filer
    _filed(data, status="proposed", sha256="c" * 64)
    assert W.run(data=data, suivi=suivi, today=date(2026, 10, 2)) == "nothing to add"
    c = db.connect(data / "harness.db"); c.execute("UPDATE scans SET status='filed' WHERE sha256=?", ("b" * 64,)); c.commit(); c.close()
    import os, time
    old = time.time() - 600; os.utime(suivi / "Suivi.xlsx", (old, old))                  # the safeguard waits while the file changed in the last 90 s
    assert "added 1 document(s)" in W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    assert W.run(data=data, suivi=suivi, today=date(2026, 10, 2)) == "nothing to add"


def test_income_documents_say_received_and_the_pause_switch_covers_documents_too(world):
    data, suivi = world
    _filed(data, doc_type="invoice_issued", folder="Income", supplier="Gablinger", number=None, amount=20000.0, paid_date="2026-01-17", proposed_name="Gablinger_Facture_CHF20000.00_recu_17.01.2026.pdf")
    (data / SE.STATE_DIR).mkdir(parents=True, exist_ok=True); (data / SE.STATE_DIR / SE.DISABLED_FILE).touch()
    assert W.run(data=data, suivi=suivi, today=date(2026, 10, 2)).startswith("paused")
    (data / SE.STATE_DIR / SE.DISABLED_FILE).unlink()
    assert "added 4 note(s) and 1 document(s)" in W.run(data=data, suivi=suivi, today=date(2026, 10, 2))
    row = next(r for r in rows_of(suivi / "Suivi.xlsx") if r[4] == "document")
    assert "Filed invoice issued by ICT4Peace: Gablinger, CHF 20000.00, received 17.01.2026 (in 2026/Income)" == row[5]
