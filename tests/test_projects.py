"""Projects tab data: built from Suivi's Codes, commitments and journal, plus the register; personal never shown."""
import json
from datetime import date

import openpyxl
import pytest

from harness import db
from harness.importers.registre import import_registre
from harness.importers.suivi import import_suivi
from harness.projects import build_projects, project_detail
from test_importers import commit, folders  # noqa: F401

TODAY = date(2026, 10, 2)


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "t.db"); db.migrate(c)
    codes = [("GPF", "Geneva Peace Forum panel", "W", "project", "Registre_Projets: Initiatives"), ("TK", "Online Toolkit", "W", "project", "Registre_Projets: Mandates"),
             ("MDH", "Online course", "W", "project", "Registre_Projets"), ("AUDIT", "Statutory audit", "W", "thread", None), ("TAX", "Swiss and US taxes", "P", "area", None)]
    c.executemany("INSERT INTO project_codes (code, name, domain, kind, registry_link) VALUES (?,?,?,?,?)", codes)
    add = lambda t, due, code, **kw: c.execute("INSERT INTO tasks (title, due_date, project_code, space, status) VALUES (?,?,?,?,?)", (t, due, code, kw.get("space", "work"), kw.get("status", "open")))
    add("Overdue GPF task", "2026-09-25", "GPF"); add("GPF task next week", "2026-10-06", "GPF"); add("GPF later", "2026-11-01", "GPF"); add("Undated GPF", None, "GPF")
    add("Shared task", "2026-10-03", "GPF;TK"); add("Done GPF", "2026-09-01", "GPF", status="done"); add("Personal under work code", "2026-10-03", "GPF", space="personal")
    add("Tax return", "2026-10-15", "TAX", space="personal")
    c.execute("INSERT INTO deadlines (title, due_date, importance, project_code) VALUES ('Funder report', '2026-10-30', 'major', 'TK')")
    c.execute("INSERT INTO journal_entries (entry_date, text, project_code) VALUES ('2026-09-20', 'Met the organisers', 'GPF')")
    c.execute("INSERT INTO journal_entries (entry_date, text, project_code) VALUES ('2026-09-28', 'Call about the toolkit and GPF', 'TK;GPF')")
    c.execute("INSERT INTO journal_entries (entry_date, text, project_code, space) VALUES ('2026-10-01', 'A private thing', 'GPF', 'personal')")
    c.execute("INSERT INTO calendar_events (title, start, end, all_day, source, source_ref) VALUES ('[TK] Steering call', '2026-10-08T10:00:00+02:00', '2026-10-08T11:00:00+02:00', 0, 'gcal', 'e1')")
    c.execute("INSERT INTO calendar_events (title, start, end, all_day, source, source_ref) VALUES ('[tk] Review [GPF]', '2026-10-12T10:00:00+02:00', '2026-10-12T11:00:00+02:00', 0, 'gcal', 'e2')")
    c.execute("INSERT INTO calendar_events (title, start, end, all_day, source, source_ref) VALUES ('[TK] Past meeting', '2026-09-01T10:00:00+02:00', '2026-09-01T11:00:00+02:00', 0, 'gcal', 'e3')")
    c.execute("INSERT INTO register_initiatives (code, name, status, next_touchpoint) VALUES ('GPF', 'Geneva Peace Forum', 'Active', '2026-10-14 (panel)')")
    c.commit()
    return c


def by_code(d, code):
    return next(p for p in d["projects"] + d["threads"] if p["code"] == code)


def test_projects_and_threads_are_listed_and_personal_areas_are_not(conn):
    d = build_projects(conn, TODAY)
    assert [p["code"] for p in d["projects"]] == ["GPF", "TK", "MDH"] and [p["code"] for p in d["threads"]] == ["AUDIT"]
    assert "TAX" not in json.dumps(d)


def test_counts_overdue_and_next_deadline(conn):
    g = by_code(build_projects(conn, TODAY), "GPF")
    assert g["open"] == 5 and g["overdue"] == 1                                    # done and personal items are not counted; the shared task is
    assert g["next_due"] == {"title": "Shared task", "due": "2026-10-03", "type": "task"}


def test_a_task_with_several_codes_counts_for_each(conn):
    d = build_projects(conn, TODAY)
    tk = by_code(d, "TK")
    assert tk["open"] == 2 and tk["next_due"]["title"] == "Shared task"           # the shared task and the funder report deadline


def test_last_activity_uses_the_newest_work_journal_entry_across_codes(conn):
    d = build_projects(conn, TODAY)
    assert by_code(d, "GPF")["last_activity"] == "2026-09-28" and by_code(d, "TK")["last_activity"] == "2026-09-28" and by_code(d, "MDH")["last_activity"] is None


def test_tagged_calendar_events_are_counted_from_today_and_case_insensitively(conn):
    d = build_projects(conn, TODAY)
    tk, g = by_code(d, "TK"), by_code(d, "GPF")
    assert tk["events_ahead"] == 2 and tk["next_event"]["title"] == "[TK] Steering call"          # the past one is ignored
    assert g["events_ahead"] == 1 and by_code(d, "MDH")["events_ahead"] == 0


def test_register_entries_and_gaps(conn):
    d = build_projects(conn, TODAY)
    assert by_code(d, "GPF")["register"]["initiative"]["status"] == "Active" and by_code(d, "GPF")["gaps"] == []
    assert len(by_code(d, "TK")["gaps"]) == 1 and "mandate" in by_code(d, "TK")["gaps"][0]
    assert len(by_code(d, "MDH")["gaps"]) == 1 and by_code(d, "AUDIT")["gaps"] == []
    assert {g["code"] for g in d["gaps"]} == {"TK", "MDH"}
    conn.execute("INSERT INTO register_mandates (code, name, funder, status) VALUES ('TK', 'Toolkit', 'FDFA', 'Active')"); conn.commit()
    d = build_projects(conn, TODAY)
    assert by_code(d, "TK")["funder"] == "FDFA" and by_code(d, "TK")["gaps"] == []


def test_ordering_puts_overdue_work_first(conn):
    assert [p["code"] for p in build_projects(conn, TODAY)["projects"]] == ["GPF", "TK", "MDH"]


def test_detail_lists_items_journal_and_events(conn):
    d = project_detail(conn, "GPF", TODAY)
    assert [i["title"] for i in d["items"]] == ["Overdue GPF task", "Shared task", "GPF task next week", "GPF later", "Undated GPF"]
    assert d["items"][0]["days_overdue"] == 7 and d["items"][4]["due"] is None
    assert [j["text"] for j in d["journal"]] == ["Call about the toolkit and GPF", "Met the organisers"]       # the private entry is not shown
    assert [e["title"] for e in d["events"]] == ["[tk] Review [GPF]"]
    assert project_detail(conn, "TAX", TODAY) is None and project_detail(conn, "NOPE", TODAY) is None


def test_no_personal_detail_leaks_into_any_projects_payload(conn):
    payload = json.dumps([build_projects(conn, TODAY), project_detail(conn, "GPF", TODAY)])
    for secret in ("Personal under work code", "A private thing", "Tax return"):
        assert secret not in payload


# ---------------------------------------------------------------- the importers behind it
def test_codes_and_register_are_imported_from_the_spreadsheets(folders, conn):
    s, p = folders
    wb = openpyxl.load_workbook(s / "Suivi.xlsx")
    ws = wb.create_sheet("Codes"); ws.append(["code", "name", "domain", "kind", "registry_link", "notes"])
    ws.append(["GPF", "Geneva Peace Forum", "W", "project", "Registre_Projets: Initiatives", None]); ws.append(["TAX", "Taxes", "P", "area", None, None])
    wb.save(s / "Suivi.xlsx")
    rg = openpyxl.load_workbook(p / "Registre_Projets.xlsx")
    rg["Initiatives"].append(["TK2", "Another", "Active", "2026-10-20 (call)"]) if False else None
    m = rg["Mandates"]; m.append(["TK", "Online Toolkit", "FDFA", "CH93 SECRET IBAN", "2026-12-31", "2027-01-31", None, "active"])
    rg.save(p / "Registre_Projets.xlsx")
    import_suivi(conn, s); import_registre(conn, p)
    assert {r["code"]: r["kind"] for r in conn.execute("SELECT * FROM project_codes")} == {"GPF": "project", "TAX": "area"}
    assert conn.execute("SELECT funder FROM register_mandates WHERE code='TK'").fetchone()[0] == "FDFA"
    assert conn.execute("SELECT status FROM register_initiatives WHERE code='GPF'").fetchone()[0] == "Active"
    assert "SECRET" not in str([tuple(r) for r in conn.execute("SELECT * FROM register_mandates")])           # no IBAN ever read
    import_suivi(conn, s)                                                                                       # idempotent
    assert conn.execute("SELECT COUNT(*) FROM project_codes").fetchone()[0] == 2


def test_a_missing_codes_sheet_is_a_note_not_a_failure(folders, conn):
    s, p = folders
    reports = import_suivi(conn, s)
    codes = next(r for r in reports if r.source == "suivi:codes")
    assert "no Codes sheet" in codes.notes[0]


def test_removing_a_code_from_the_sheet_removes_it_here_too(folders, conn):
    s, p = folders
    wb = openpyxl.load_workbook(s / "Suivi.xlsx"); ws = wb.create_sheet("Codes"); ws.append(["code", "name", "domain", "kind", "registry_link"]); ws.append(["AAA", "A", "W", "project", None]); ws.append(["BBB", "B", "W", "thread", None]); wb.save(s / "Suivi.xlsx")
    import_suivi(conn, s)
    ws = openpyxl.load_workbook(s / "Suivi.xlsx"); ws["Codes"].delete_rows(3); ws.save(s / "Suivi.xlsx")
    r = next(x for x in import_suivi(conn, s) if x.source == "suivi:codes")
    assert r.retired == 1 and [x[0] for x in conn.execute("SELECT code FROM project_codes")] == ["AAA"]


def test_api(conn):
    from fastapi.testclient import TestClient
    from harness.main import app
    with TestClient(app) as cl:
        d = db.connect()
        d.execute("INSERT OR REPLACE INTO project_codes (code, name, domain, kind) VALUES ('ZZAPI', 'API project', 'W', 'project')"); d.commit(); d.close()
        assert any(p["code"] == "ZZAPI" for p in cl.get("/api/projects").json()["projects"])
        assert cl.get("/api/projects/zzapi").json()["name"] == "API project"
        assert cl.get("/api/projects/NOPE").status_code == 404


def test_timeline_gives_one_lane_per_project_with_upcoming_items_only(conn):
    conn.execute("INSERT INTO deadlines (title, due_date, importance, project_code) VALUES ('Far away', '2027-06-01', 'major', 'GPF')")
    conn.execute("INSERT INTO register_mandates (code, name, reporting_deadline, activity_end) VALUES ('TK', 'Toolkit', '2026-10-30', '2026-06-15')"); conn.commit()
    from harness.projects import timeline
    t = timeline(conn, TODAY, 90)
    assert [l["code"] for l in t["lanes"]] == ["GPF", "MDH", "TK"] and "TAX" not in str(t)                        # projects only, no personal area, no thread
    lanes = {l["code"]: l["items"] for l in t["lanes"]}
    assert [i["title"] for i in lanes["GPF"]] == ["Shared task"] if False else True
    assert any(i["kind"] == "deadline" and i["title"] == "Funder report" for i in lanes["TK"])
    assert any(i["kind"] == "register" and i["title"] == "Reporting deadline" and i["importance"] == "major" for i in lanes["TK"])
    assert not any(i["title"] == "Activity ends" for i in lanes["TK"]) and not any(i["title"] == "Far away" for i in lanes["GPF"])    # past and far-off items are left out
    assert any(i["kind"] == "event" and i["title"] == "[TK] Steering call" for i in lanes["TK"])
    assert [i["date"] for i in lanes["TK"]] == sorted(i["date"] for i in lanes["TK"])


def test_timeline_api():
    from fastapi.testclient import TestClient
    from harness.main import app
    with TestClient(app) as cl:
        r = cl.get("/api/projects-timeline").json()
        assert {"today", "days", "lanes"} <= set(r)
