"""The register writer: append-only, backed up, verified, reversible. Synthetic workbooks only (shaped like the real register)."""
import json
import os
import sys
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import register_writer as RW  # noqa: E402

COLS = {"Mandates": ["code", "name", "funder", "contract no.", "order no.", "signature", "activity start", "activity end", "reporting deadline", "interim reports", "max contribution", "ceiling %", "own contribution required",
                     "dedicated account", "IBAN", "visibility requirement", "visibility satisfied where", "audit financial year", "substantive officer", "processing officer", "status", "close-out state", "notes"],
        "Instalments": ["code", "instalment", "amount", "expected date", "received date", "account", "reference"],
        "Budget": ["code", "budget line", "phase", "funder amount", "own amount", "total"],
        "Initiatives": ["code", "name", "strand", "counterpart", "status", "next touchpoint", "notes"]}


def _cell(ref, text, s):
    return f'<c r="{ref}" s="{s}" t="inlineStr"><is><t>{text}</t></is></c>'


def sheet_xml(sheet, existing=()):
    L = lambda n: chr(64 + n) if n <= 26 else "A" + chr(64 + n - 26)
    head = "".join(_cell(f"{L(i)}1", c, 1) for i, c in enumerate(COLS[sheet], 1))
    rows = f'<row r="1" ht="26" customHeight="1">{head}</row>'
    for r, vals in enumerate(existing, 2):
        rows += f'<row r="{r}">' + "".join(_cell(f"{L(i)}{r}", v, 2) for i, v in enumerate(vals, 1)) + "</row>"
    last = 1 + len(existing)
    return (f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetPr><outlinePr summaryBelow="1" summaryRight="1"/><pageSetUpPr/></sheetPr><dimension ref="A1:{L(len(COLS[sheet]))}{last}"/>'
            f'<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews><sheetFormatPr baseColWidth="8" defaultRowHeight="15"/>'
            f'<cols><col width="12" customWidth="1" min="1" max="1"/></cols><sheetData>{rows}</sheetData><pageMargins left="0.75" right="0.75" top="1" bottom="1" header="0.5" footer="0.5"/></worksheet>')


def make_register(path, initiatives=(("GPF", "Geneva Peace Forum", "AI for peace", "Globethics", "Active", "2026-10-14", "notes"),)):
    names = list(COLS)
    wb = "".join(f'<sheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" name="{n}" sheetId="{i}" state="visible" r:id="rId{i}"/>' for i, n in enumerate(names, 1))
    rels = "".join(f'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet{i}.xml" Id="rId{i}"/>' for i in range(1, 5))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        ct = "".join(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' for i in range(1, 5))
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                   f'<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>{ct}</Types>')
        z.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="/xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", f'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheets>{wb}</sheets><definedNames/></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>')
        z.writestr("xl/styles.xml", '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><fonts count="1"><font><sz val="11"/></font></fonts><fills count="2"><fill><patternFill/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
                   '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
                   '<cellXfs count="3"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs></styleSheet>')
        z.writestr("xl/theme/theme1.xml", "<theme/>")
        for i, n in enumerate(names, 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", sheet_xml(n, initiatives if n == "Initiatives" else ()))
    old = time.time() - 3600
    os.utime(path, (old, old))


ROWS = {"Mandates": [{"code": "TK", "name": "Toolkit & <FDFA>", "funder": "Swiss FDFA (DFAE)", "contract no.": "81086897", "max contribution": 49900, "ceiling %": "60%", "notes": "Line one\nLine \"two\" & more"}],
        "Instalments": [{"code": "TK", "instalment": "1st instalment (50%)", "amount": 24950, "received date": "2026-01-06"}, {"code": "TK", "instalment": "Final", "amount": 24878.9, "received date": "2026-09-22"}],
        "Budget": [{"code": "TK", "budget line": "Content review & update", "phase": "Phase 1", "funder amount": 3000}],
        "Initiatives": [{"code": "FAGI", "name": "FAGI consortium bid", "counterpart": "Globethics (lead)", "status": "Bid"}, {"code": "GPF", "name": "should be skipped"}]}


@pytest.fixture
def reg(tmp_path):
    p = tmp_path / "Registre_Projets.xlsx"
    make_register(p)
    return p


def entries(p):
    z = zipfile.ZipFile(p)
    return {n: z.read(n) for n in z.namelist()}


def test_rows_are_appended_and_everything_else_is_byte_for_byte_unchanged(reg, tmp_path):
    before = entries(reg)
    r = RW.apply(reg, ROWS, backups=tmp_path / "bk")
    assert r["written"] and len(r["added"]) == 5 and r["skipped"] == ["Initiatives: GPF is already there"]
    after = entries(reg)
    assert set(after) == set(before)
    for name in before:
        if name not in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml", "xl/worksheets/sheet4.xml"):
            assert after[name] == before[name], name
    import re
    strip = lambda x: re.sub(r"<dimension[^>]*/>", "", x)
    old4, new4 = strip(before["xl/worksheets/sheet4.xml"].decode()), strip(after["xl/worksheets/sheet4.xml"].decode())
    old_data = old4.split("</sheetData>")[0]
    assert new4.startswith(old_data) and old4.split("</sheetData>")[1] == new4.split("</sheetData>")[1]                         # the old rows are an exact prefix; what follows the data is unchanged
    assert 'r="3"' in new4 and "FAGI consortium bid" in new4 and 'ref="A1:G3"' in after["xl/worksheets/sheet4.xml"].decode().replace("<dimension ", " ")


def test_the_file_opens_in_openpyxl_with_the_new_rows_in_the_right_columns_and_real_numbers(reg, tmp_path):
    RW.apply(reg, ROWS, backups=tmp_path / "bk")
    wb = openpyxl.load_workbook(reg)
    m = wb["Mandates"]
    row = {m.cell(1, c).value: m.cell(2, c).value for c in range(1, 24)}
    assert (row["code"], row["name"], row["funder"], row["contract no."], row["max contribution"], row["ceiling %"]) == ("TK", "Toolkit & <FDFA>", "Swiss FDFA (DFAE)", "81086897", 49900, "60%")
    assert row["notes"] == 'Line one\nLine "two" & more' and row["IBAN"] is None
    i = wb["Instalments"]
    assert [(i.cell(r, 2).value, i.cell(r, 3).value, i.cell(r, 5).value) for r in (2, 3)] == [("1st instalment (50%)", 24950, "2026-01-06"), ("Final", 24878.9, "2026-09-22")]
    assert [c.value for c in wb["Initiatives"][3]][:2] == ["FAGI", "FAGI consortium bid"] and wb["Initiatives"].max_row == 3
    assert wb["Budget"].cell(2, 2).value == "Content review & update"


def test_the_harness_importer_reads_the_written_register(reg, tmp_path):
    from harness import db
    from harness.importers.registre import import_registre
    RW.apply(reg, ROWS, backups=tmp_path / "bk")
    c = db.connect(tmp_path / "t.db"); db.migrate(c)
    import_registre(c, reg.parent)
    assert c.execute("SELECT funder FROM register_mandates WHERE code='TK'").fetchone()[0] == "Swiss FDFA (DFAE)"
    assert c.execute("SELECT budget_line FROM register_budget WHERE code='TK'").fetchone()[0] == "Content review & update"
    assert c.execute("SELECT name FROM register_initiatives WHERE code='FAGI'").fetchone()[0] == "FAGI consortium bid"


def test_a_backup_identical_to_the_original_is_made_first(reg, tmp_path):
    original = reg.read_bytes()
    r = RW.apply(reg, ROWS, backups=tmp_path / "bk")
    assert Path(r["backup"]).read_bytes() == original


def test_running_twice_adds_nothing_the_second_time(reg, tmp_path):
    RW.apply(reg, ROWS, backups=tmp_path / "bk")
    after = reg.read_bytes()
    r = RW.apply(reg, ROWS, backups=tmp_path / "bk")
    assert r["written"] is False and len(r["skipped"]) == 6 and reg.read_bytes() == after


def test_it_refuses_while_excel_has_the_file_open_or_it_changed_a_moment_ago(reg, tmp_path):
    before = reg.read_bytes()
    (reg.parent / f"~${reg.name}").write_bytes(b"lock")
    with pytest.raises(RW.RegisterError, match="Excel"):
        RW.apply(reg, ROWS, backups=tmp_path / "bk")
    (reg.parent / f"~${reg.name}").unlink()
    os.utime(reg, None)
    with pytest.raises(RW.RegisterError, match="90 seconds"):
        RW.apply(reg, ROWS, backups=tmp_path / "bk")
    assert reg.read_bytes() == before and not (tmp_path / "bk").exists()


def test_unknown_sheets_columns_or_rows_without_a_key_are_refused_and_nothing_is_changed(reg, tmp_path):
    before = reg.read_bytes()
    for bad in ({"Nope": [{"code": "X"}]}, {"Mandates": [{"code": "X", "colour": "red"}]}, {"Mandates": [{"name": "no code"}]}, {"Budget": [{"code": "X"}]}):
        with pytest.raises(RW.RegisterError):
            RW.apply(reg, bad, backups=tmp_path / "bk")
    assert reg.read_bytes() == before


def test_if_the_final_check_fails_her_original_is_restored_exactly(reg, tmp_path, monkeypatch):
    original = reg.read_bytes()
    calls = {"n": 0}
    real = RW.verify
    def flaky(*a, **k):
        calls["n"] += 1
        if calls["n"] == 3:                                     # the check of the file after it replaced the original
            raise RW.RegisterError("simulated trouble")
        return real(*a, **k)
    monkeypatch.setattr(RW, "verify", flaky)
    with pytest.raises(RW.RegisterError, match="nothing was changed"):
        RW.apply(reg, ROWS, backups=tmp_path / "bk")
    assert reg.read_bytes() == original and not list(reg.parent.glob(".*.new"))


def test_a_workbook_saved_by_excel_style_tools_with_shared_strings_works_too(tmp_path):
    p = tmp_path / "Registre_Projets.xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for name, cols in COLS.items():
        wb.create_sheet(name).append(cols)
    wb["Initiatives"].append(["GPF", "Geneva Peace Forum", None, None, "Active"])
    wb.save(p)
    os.utime(p, (time.time() - 3600,) * 2)
    r = RW.apply(p, ROWS, backups=tmp_path / "bk")
    assert len(r["added"]) == 5 and r["skipped"] == ["Initiatives: GPF is already there"]
    w = openpyxl.load_workbook(p)
    assert w["Initiatives"].max_row == 3 and w["Initiatives"].cell(2, 2).value == "Geneva Peace Forum" and w["Mandates"].cell(2, 1).value == "TK"


def test_preview_changes_nothing(reg, tmp_path):
    before = reg.read_bytes()
    r = RW.apply(reg, ROWS, backups=tmp_path / "bk", dry_run=True)
    assert r["written"] is False and len(r["added"]) == 5 and reg.read_bytes() == before and not (tmp_path / "bk").exists()


def test_source_never_deletes_or_rewrites_the_register_other_than_the_verified_swap():
    src = Path(RW.__file__).read_text()
    for banned in ("os.remove", "shutil.rmtree", ".truncate(", 'open(path, "w'):
        assert banned not in src, banned
    assert src.count("os.replace(tmp, path)") == 1 and src.count("shutil.copy2(backup, path)") == 1
