"""Editing title and due date in Today & overdue, and keeping those edits across re-imports."""
from datetime import datetime

import openpyxl
import pytest

from harness import db
from harness.config import TZ
from harness.importers.run import run_all
from harness.items import EditRefused, edit_item
from harness.today import build_today
from test_importers import commit, folders  # noqa: F401  (synthetic Suivi + register fixtures)

NOW = datetime(2026, 10, 1, 9, 0, tzinfo=TZ)


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "t.db"); db.migrate(c)
    return c


def row(c, table, id): return dict(c.execute(f"SELECT * FROM {table} WHERE id=?", (id,)).fetchone())


def test_edit_a_manual_task_title_and_date(conn):
    conn.execute("INSERT INTO tasks (title, due_date) VALUES ('Old title', '2026-10-01')"); conn.commit()
    assert edit_item(conn, "task", 1, title="  New   title ", due="2026-10-09", set_due=True)
    r = row(conn, "tasks", 1)
    assert r["title"] == "New title" and r["due_date"] == "2026-10-09"
    assert conn.execute("SELECT COUNT(*) FROM item_overrides").fetchone()[0] == 0       # nothing to protect: not imported


def test_a_date_can_be_cleared_on_tasks_and_waiting_on_but_not_deadlines(conn):
    conn.execute("INSERT INTO tasks (title, due_date) VALUES ('T', '2026-10-01')")
    conn.execute("INSERT INTO waiting_on (description, since_date, remind_on) VALUES ('W', '2026-09-01', '2026-10-01')")
    conn.execute("INSERT INTO deadlines (title, due_date) VALUES ('D', '2026-10-05')"); conn.commit()
    assert edit_item(conn, "task", 1, due=None, set_due=True) and row(conn, "tasks", 1)["due_date"] is None
    assert edit_item(conn, "waiting_on", 1, title="W2", due=None, set_due=True)
    r = row(conn, "waiting_on", 1); assert r["description"] == "W2" and r["remind_on"] is None
    with pytest.raises(EditRefused, match="needs a date"):
        edit_item(conn, "deadline", 1, due="", set_due=True)


@pytest.mark.parametrize("kw,msg", [({"title": "   "}, "empty"), ({"title": "x" * 301}, "too long"), ({"due": "31/10/2026", "set_due": True}, "not valid"),
                                    ({"due": "2026-02-30", "set_due": True}, "not valid")])
def test_bad_edits_are_refused_and_change_nothing(conn, kw, msg):
    conn.execute("INSERT INTO tasks (title, due_date) VALUES ('Keep', '2026-10-01')"); conn.commit()
    with pytest.raises(EditRefused, match=msg):
        edit_item(conn, "task", 1, **kw)
    assert row(conn, "tasks", 1)["title"] == "Keep" and row(conn, "tasks", 1)["due_date"] == "2026-10-01"


def test_done_items_must_be_reopened_first_and_missing_items_are_reported(conn):
    conn.execute("INSERT INTO tasks (title, status) VALUES ('Done one', 'done')"); conn.commit()
    with pytest.raises(EditRefused, match="Reopen"):
        edit_item(conn, "task", 1, title="x")
    assert edit_item(conn, "task", 99, title="x") is False


def test_edit_moves_the_item_between_lists_and_moves_the_deadline_on_the_lake(conn):
    conn.execute("INSERT INTO tasks (title, due_date) VALUES ('Move me', '2026-09-25')")
    conn.execute("INSERT INTO deadlines (title, due_date, importance) VALUES ('Report', '2026-10-05', 'major')"); conn.commit()
    assert [i["title"] for i in build_today(conn, NOW)["today_items"]] == ["Move me"]
    edit_item(conn, "task", 1, due="2026-10-20", set_due=True)
    t = build_today(conn, NOW)
    assert "Move me" not in [i["title"] for i in t["today_items"] + t["upcoming_items"]]
    assert [i["title"] for i in t["later_items"]] == ["Move me"]                         # not lost: listed under Later
    edit_item(conn, "deadline", 1, due="2026-10-12", set_due=True)
    t = build_today(conn, NOW)
    assert t["lake"][0]["due"] == "2026-10-12" and t["ticks"] == ["2026-10-09"]          # D-3 warning follows the new date (D-14 already passed)


def test_hand_edits_to_suivi_items_survive_reimport_and_track_the_source(folders, conn):
    s, p = folders
    run_all(conn, s, p)
    tid = conn.execute("SELECT id FROM tasks WHERE source_ref='C-1'").fetchone()[0]
    edit_item(conn, "task", tid, title="My own wording", due="2026-10-15", set_due=True)
    run_all(conn, s, p)                                                                  # re-import: must not undo it
    r = row(conn, "tasks", tid)
    assert r["title"] == "My own wording" and r["due_date"] == "2026-10-15"
    t = build_today(conn, NOW)
    it = next(i for i in t["upcoming_items"] + t["later_items"] if i["type"] == "task" and i["id"] == tid)
    assert it["edited"]["title"] == "Soft task" and it["edited"]["due"] == "2026-10-08" and it["edited"]["title_changed"] and it["edited"]["due_changed"]

    wb = openpyxl.load_workbook(s / "Suivi.xlsx"); ws = wb["Commitments"]
    ws["G2"] = "Soft task (changed in Suivi)"; ws["H2"] = "2026-10-30"; wb.save(s / "Suivi.xlsx")
    run_all(conn, s, p)
    assert row(conn, "tasks", tid)["title"] == "My own wording"                          # still mine
    it = next(i for i in build_today(conn, NOW)["later_items"] if i["type"] == "task" and i["id"] == tid)
    assert it["edited"]["title"] == "Soft task (changed in Suivi)" and it["edited"]["due"] == "2026-10-30"   # but I can see what Suivi now says

    edit_item(conn, "task", tid, reset=True)
    r = row(conn, "tasks", tid)
    assert r["title"] == "Soft task (changed in Suivi)" and r["due_date"] == "2026-10-30"
    assert conn.execute("SELECT COUNT(*) FROM item_overrides").fetchone()[0] == 0


def test_editing_back_to_the_sources_value_clears_the_override(folders, conn):
    s, p = folders; run_all(conn, s, p)
    tid = conn.execute("SELECT id FROM tasks WHERE source_ref='C-1'").fetchone()[0]
    edit_item(conn, "task", tid, title="Changed")
    assert conn.execute("SELECT COUNT(*) FROM item_overrides").fetchone()[0] == 1
    edit_item(conn, "task", tid, title="Soft task")
    assert conn.execute("SELECT COUNT(*) FROM item_overrides").fetchone()[0] == 0


def test_edited_waiting_on_and_deadline_from_suivi_are_protected_too(folders, conn):
    s, p = folders; run_all(conn, s, p)
    wid = conn.execute("SELECT id FROM waiting_on").fetchone()[0]
    did = conn.execute("SELECT id FROM deadlines WHERE source_ref='C-2'").fetchone()[0]
    edit_item(conn, "waiting_on", wid, due="2026-11-05", set_due=True); edit_item(conn, "deadline", did, due="2026-11-20", set_due=True)
    run_all(conn, s, p)
    assert row(conn, "waiting_on", wid)["remind_on"] == "2026-11-05" and row(conn, "deadlines", did)["due_date"] == "2026-11-20"


def test_api_edit_endpoint():
    from fastapi.testclient import TestClient
    from harness.main import app
    with TestClient(app) as cl:
        c = db.connect()
        c.execute("INSERT INTO tasks (title, due_date) VALUES ('API edit me', '2000-01-01')"); c.commit()
        tid = c.execute("SELECT id FROM tasks WHERE title='API edit me'").fetchone()[0]; c.close()
        assert cl.patch(f"/api/items/task/{tid}", json={"title": "Edited by API", "due": "2000-02-02"}).json() == {"ok": True}
        assert cl.patch(f"/api/items/task/{tid}", json={"title": "Only title"}).json() == {"ok": True}
        c = db.connect(); r = dict(c.execute("SELECT title, due_date FROM tasks WHERE id=?", (tid,)).fetchone()); c.close()
        assert r == {"title": "Only title", "due_date": "2000-02-02"}                    # 'due' omitted = unchanged
        assert cl.patch(f"/api/items/task/{tid}", json={"due": None}).json() == {"ok": True}
        c = db.connect(); assert c.execute("SELECT due_date FROM tasks WHERE id=?", (tid,)).fetchone()[0] is None; c.close()
        assert cl.patch(f"/api/items/task/{tid}", json={"title": ""}).status_code == 422
        assert cl.patch("/api/items/task/999999", json={"title": "x"}).status_code == 404
        assert cl.patch(f"/api/items/email/{tid}", json={"title": "x"}).status_code == 404
