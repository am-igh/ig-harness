"""Importer tests use synthetic spreadsheets built on the fly (rule 8)."""
import openpyxl
import pytest

from harness import db
from harness.importers.run import run_all

COMMIT_COLS = ["id", "created", "domain", "code", "direction", "counterparty", "what",
               "due", "weight", "status", "reminder", "origin", "closed_on", "close_evidence"]
JOURNAL_COLS = ["id", "date", "domain", "code", "type", "summary", "people", "evidence",
                "hours", "source", "status", "entered_on"]


def _book(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)
    wb.save(path)


def commit(id, what, **kw):
    d = dict(id=id, created="2026-09-01", domain="W", code="PX", direction="owe", counterparty="Partner A",
             what=what, due=None, weight="soft", status="open", reminder=None, origin=None,
             closed_on=None, close_evidence=None)
    d.update(kw)
    return [d[c] for c in COMMIT_COLS]


@pytest.fixture
def folders(tmp_path):
    s, p = tmp_path / "suivi", tmp_path / "projets"
    s.mkdir(); p.mkdir()
    _book(s / "Suivi.xlsx", [
        ("Commitments", COMMIT_COLS, [
            commit("C-1", "Soft task", due="2026-10-08"),
            commit("C-2", "Hard deadline", due="2026-10-20", weight="hard"),
            commit("C-3", "Major deadline", due="2026-11-01", weight="major"),
            commit("C-4", "Partner owes us a draft", direction="owed", due="2026-10-10"),
            commit("C-5", "Personal tax form", domain="P", due="2026-10-30"),
            commit("C-6", "Only a proposal", status="proposed"),
            commit("C-7", "Already done", status="done", closed_on="2026-09-20"),
            commit("C-8", "Odd date", due="soon"),
        ]),
        ("Journal", JOURNAL_COLS, [
            ["J-1", "2026-09-30", "W", "PX", "email", "Wrote to funder", None, None, None, None, "confirmed", "2026-09-30"],
            ["J-2", "2026-09-29", "P", None, "admin", "Personal note", None, None, None, None, "confirmed", "2026-09-29"],
        ]),
    ])
    _book(p / "Registre_Projets.xlsx", [
        ("Mandates", ["code", "name", "funder", "IBAN", "activity end", "reporting deadline",
                      "interim reports", "status"], [
            ["PX", "Project X", "Funder", "CH00 SECRET", "2026-12-31", "2027-01-31",
             "2026-10-31; 2026-11-30", "active"],
            ["PY", "Old project", "F", "CH11", "2025-12-31", "2026-01-31", None, "closed"],
        ]),
        ("Instalments", ["code", "instalment", "amount", "expected date", "received date"], [
            ["PX", "2", 1000, "2026-11-15", None],
            ["PX", "1", 1000, "2026-09-15", "2026-09-16"],
        ]),
        ("Initiatives", ["code", "name", "status", "next touchpoint"], [
            ["GPF", "Forum", "Active", "2026-10-14 (panel)"],
            ["ZZZ", "Vague", "Active", "sometime"],
        ]),
    ])
    return s, p


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "t.db")
    db.migrate(c)
    return c


def q(conn, sql):
    return [dict(r) for r in conn.execute(sql)]


def test_suivi_mapping(conn, folders):
    run_all(conn, *folders)
    tasks = {r["source_ref"]: r for r in q(conn, "SELECT * FROM tasks")}
    assert set(tasks) == {"C-1", "C-5", "C-7", "C-8"}  # proposed skipped; hard/major are deadlines
    assert tasks["C-7"]["status"] == "done" and tasks["C-7"]["done_at"] == "2026-09-20"
    assert tasks["C-8"]["due_date"] is None  # unparseable date left empty, not guessed
    dl = {r["source_ref"]: r for r in q(conn, "SELECT * FROM deadlines WHERE source='suivi'")}
    assert dl["C-2"]["importance"] == "normal" and dl["C-3"]["importance"] == "major"
    w = q(conn, "SELECT * FROM waiting_on")
    assert len(w) == 1 and w[0]["person"] == "Partner A" and w[0]["remind_on"] == "2026-10-10"


def test_personal_items_are_s3_personal(conn, folders):
    run_all(conn, *folders)
    t = q(conn, "SELECT space, sensitivity FROM tasks WHERE source_ref='C-5'")[0]
    j = q(conn, "SELECT space, sensitivity FROM journal_entries WHERE source_ref='J-2'")[0]
    assert t == j == {"space": "personal", "sensitivity": "S3"}


def test_registre_deadlines(conn, folders):
    run_all(conn, *folders)
    refs = {r["source_ref"]: r for r in q(conn, "SELECT * FROM deadlines WHERE source='registre'")}
    assert set(refs) == {"PX|reporting", "PX|interim|2026-10-31", "PX|interim|2026-11-30",
                         "PX|activity-end", "PX|instalment|2", "GPF|touchpoint"}
    assert refs["PX|reporting"]["importance"] == "major"
    assert refs["GPF|touchpoint"]["kind"] == "fixed"
    assert refs["GPF|touchpoint"]["title"] == "Forum: panel"


def test_no_sensitive_values_reach_database(conn, folders):
    run_all(conn, *folders)
    dump = " ".join(str(tuple(r)) for t in ("tasks", "deadlines", "waiting_on", "journal_entries")
                    for r in conn.execute(f"SELECT * FROM {t}"))
    assert "SECRET" not in dump and "1000" not in dump  # IBAN and amounts never imported


def test_import_is_idempotent(conn, folders):
    run_all(conn, *folders)
    again = run_all(conn, *folders)
    assert all(r["added"] == 0 and r["updated"] == 0 and r["retired"] == 0 for r in again)
    assert sum(r["unchanged"] for r in again) > 0


def test_local_done_survives_reimport_and_changes_are_reported(conn, folders):
    s, p = folders
    run_all(conn, *folders)
    conn.execute("UPDATE tasks SET status='done' WHERE source_ref='C-1'")
    conn.commit()
    wb = openpyxl.load_workbook(s / "Suivi.xlsx")
    ws = wb["Commitments"]
    ws["G2"] = "Soft task (renamed)"   # C-1 text
    ws.delete_rows(3)                    # C-2 vanishes from the source
    wb.save(s / "Suivi.xlsx")
    reps = {r["source"]: r for r in run_all(conn, *folders)}
    c = reps["suivi:commitments"]
    assert c["updated"] == 1 and c["retired"] == 1
    t = q(conn, "SELECT title, status FROM tasks WHERE source_ref='C-1'")[0]
    assert t == {"title": "Soft task (renamed)", "status": "done"}  # not reopened
    assert q(conn, "SELECT status FROM deadlines WHERE source_ref='C-2'")[0]["status"] == "dropped"


def test_missing_file_reports_error_without_crashing(conn, tmp_path):
    empty = tmp_path / "none"; empty.mkdir()
    reps = run_all(conn, empty, empty)
    assert len(reps) == 2 and all("error" in r for r in reps)
